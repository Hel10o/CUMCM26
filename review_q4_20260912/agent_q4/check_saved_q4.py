"""Independent read-only audit of frozen Q4 arrays; does not integrate the PDE."""
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import csv, hashlib, json, platform
import numpy as np
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parents[2] / 'q4_complete_delivery/v1'
OUT = Path(__file__).resolve().parent
z = np.load(BASE / 'output/main.npz')
event = json.loads((BASE / 'output/end_event.json').read_text(encoding='utf-8'))
main = json.loads((BASE / 'output/main.json').read_text(encoding='utf-8'))
checks = {}
def rounded(value):
    return float(Decimal(str(float(value))).quantize(Decimal('0.0001'), rounding=ROUND_HALF_UP))

t, radius, fields = z['time_s'], z['radius_m'], z['full_TC']
endidx = int(np.argmin(abs(z['event_time_s'] - event['execution_s'])))
checks['endpoint'] = {
    'saved_output_states': len(t), 'formal_data_rows': int((t > 0).sum()),
    'last_time_s': float(t[-1]), 'execution_h': float(t[-1]/3600),
    'last_max_nodal_C': float(fields[-1,:,1].max()),
    'event_last_state_max_difference': float(abs(fields[-1]-z['event_full_TC'][endidx]).max()),
    'stored_polynomial_max_minus_axis_max_abs': float(abs(z['max_C']-fields[:,0,1]).max()),
    'strict_under_frozen_1d_model': bool(fields[-1,:,1].max() < .15),
    'output_step_before_endpoint_max_deviation_from_60_s': float(abs(np.diff(t[:-1])-60).max()),
}
assert t[-1] == event['execution_s']
assert checks['endpoint']['strict_under_frozen_1d_model']
assert checks['endpoint']['event_last_state_max_difference'] == 0
assert np.all(np.diff(t[:-1]) == 60)

balance = z['balance']
mean_independent = fields[:,:,1] @ z['quadrature_weights']
checks['mass'] = {
    'mean_C_recomputed_max_abs_difference': float(abs(mean_independent-z['mean_C']).max()),
    'mean_C_end': float(mean_independent[-1]),
    'stored_boundary_integral_end': float(balance[-1,2]),
    'end_balance_from_saved_integral': float(mean_independent[-1]+balance[-1,2]-2.55),
    'max_abs_saved_balance': float(abs(balance[:,-1]).max()),
    'scope': 'Recomputed weighted mean using saved quadrature weights, and read saved boundary integral; no independent new flux integration or PDE run.',
}

domain = z['fixed_radius_m'][None,:] > radius[:,None]+1e-14
checks['domain'] = {
    'outside_cells': int(domain.sum()),
    'outside_all_nan': bool(np.isnan(z['sample_TC'][:,:,1][domain]).all()),
    'inside_all_finite': bool(np.isfinite(z['sample_TC'][:,:,1][~domain]).all()),
}
assert checks['domain']['outside_all_nan'] and checks['domain']['inside_all_finite']

wb = load_workbook(BASE/'output/result4.xlsx', read_only=True, data_only=True)
ws = wb['Sheet1']
rows = list(ws.values)
checks['workbook'] = {'sheet_names':wb.sheetnames, 'header':list(rows[0]), 'worksheet_rows':len(rows)}
wb.close()
formal = np.flatnonzero(t > 0)
errors=[]
checked_values=0
for row_i,i in enumerate(formal, start=1):
    row=rows[row_i]
    expected=[float(t[i]), *[None if not np.isfinite(x) else rounded(x) for x in z['sample_TC'][i,:,1]], rounded(z['surface_TC'][i,1]), None, float(radius[i]*100)]
    # Column X is a spacer; delivered radius metadata is column Y.
    if len(row)<len(expected): errors.append([row_i+1,'columns',len(row),len(expected)])
    for j,(actual,want) in enumerate(zip(row,expected)):
        checked_values+=1
        if want is None:
            if actual is not None: errors.append([row_i+1,j+1,actual,want])
        elif actual is None or abs(float(actual)-want)>1e-10:
            errors.append([row_i+1,j+1,actual,want])
assert len(rows)-1 == len(formal)
checks['workbook'].update({'checked_cells':checked_values,'mismatch_count':len(errors),'first_mismatches':errors[:10]})
assert not errors, errors[:10]

table=[]
with (BASE/'output/table6_unrounded.csv').open(encoding='utf-8', newline='') as f:
    for row in csv.DictReader(f):
        h=float(row['time_h']); idx=int(np.argmin(abs(t-h*3600)))
        values=[h,*z['sample_TC'][idx,::5,1],z['surface_TC'][idx,1],radius[idx]*100]
        actual=[float(v) if v else np.nan for v in row.values()]
        assert np.allclose(actual,values,rtol=0,atol=1e-11,equal_nan=True)
        table.append({'time_h':h,'matches_saved_array':True})
checks['table6']=table

checks['saved_run_comparisons']={}
for name in ('spectral80','spectral160','time120_bdf'):
    q=np.load(BASE/f'validation/{name}.npz')
    other=json.loads((BASE/f'validation/{name}.json').read_text(encoding='utf-8'))
    common,ia,ib=np.intersect1d(t,q['time_s'],return_indices=True)
    diff=z['sample_TC'][ia]-q['sample_TC'][ib]
    finite=np.isfinite(diff)
    mismatch=sum(rounded(a)!=rounded(b) for a,b in zip(z['sample_TC'][ia][finite],q['sample_TC'][ib][finite]))
    checks['saved_run_comparisons'][name]={
        'common_times':len(common),'max_abs_T':float(np.nanmax(abs(diff[:,:,0]))),
        'max_abs_C':float(np.nanmax(abs(diff[:,:,1]))),
        'four_decimal_mismatch_count_TC':mismatch,
        'critical_difference_s':other['event']['critical_s']-main['event']['critical_s'],
    }
    assert mismatch == 0

# This is a diagnostic of a DIFFERENT interpretation: rho(C) as true wet density.
# The submitted model explicitly calls rho(C) effective heat-property density.
wet_density = 760+90*fields[:,:,1]
implied_dry_density = wet_density/(1+fields[:,:,1])
md_ratio = (radius/.02)**2 * (implied_dry_density @ z['quadrature_weights']) / ((760+90*2.55)/(1+2.55))
imin,imax=int(md_ratio.argmin()),int(md_ratio.argmax())
checks['physical_density_interpretation_diagnostic']={
    'assumption_tested':'If appendix rho(C) were interpreted as actual current wet density, md(t)=integral rho(C)/(1+C) dV should be constant.',
    'initial_dry_mass_ratio':float(md_ratio[0]),'final_dry_mass_ratio':float(md_ratio[-1]),
    'min_ratio':float(md_ratio[imin]),'min_time_h':float(t[imin]/3600),
    'max_ratio':float(md_ratio[imax]),'max_time_h':float(t[imax]/3600),
    'meaning':'This exposes the need to justify the effective-density interpretation; it is not a numerical conservation failure of the submitted rho_d model.',
}

q=np.load(BASE/'validation/pure_shrinkage_evolution.npz')
checks['pure_shrinkage_saved_keys']=list(q.files)
checks['runtime']={'python':platform.python_version(),'numpy':np.__version__,'PDE_rerun':False}
checks['inspected_file_hashes']={name:hashlib.sha256((BASE/name).read_bytes()).hexdigest() for name in [
    'output/main.npz','output/main.json','output/end_event.json','output/result4.xlsx','output/table6_unrounded.csv',
    '第四问论文正文.md','第四问分析与答案.md','第四问模型与文献采用.md']}
(OUT/'saved_q4_audit.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(checks,ensure_ascii=False,indent=2))
