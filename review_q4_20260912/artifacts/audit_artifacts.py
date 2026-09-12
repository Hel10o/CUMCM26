"""Independent, read-only delivery/array/workbook audit; no PDE solve or authoring.

Run with the bundled Python (NumPy, openpyxl and lxml). Outputs are confined to
this script's directory. All supplied delivery files are opened read-only.
"""
from __future__ import annotations

import csv
import hashlib
import json
import math
import platform
import re
import sys
import time
import zipfile
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import numpy as np
import openpyxl
from lxml import etree as ET

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
DELIVERY = ROOT / "q4_complete_delivery/v1"
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def sha(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for block in iter(lambda: f.read(4 * 1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def readj(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def r4(value):
    if value is None or not math.isfinite(value):
        return None
    return float(Decimal(str(float(value))).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP))


def round_array(values):
    scaled = np.abs(values) * 10000
    rounded = np.sign(values) * np.floor(scaled + 0.5) / 10000
    ties = np.abs(scaled - np.floor(scaled) - 0.5) <= 1e-8
    for index in zip(*np.where(ties)):
        rounded[index] = r4(float(values[index]))
    return rounded


def array_check(actual, expected):
    a, b = np.asarray(actual, dtype=float), np.asarray(expected, dtype=float)
    if a.shape != b.shape:
        return {"passed": False, "actual_shape": list(a.shape), "expected_shape": list(b.shape)}
    both = np.isfinite(a) & np.isfinite(b)
    nonfinite_mismatches = int(np.count_nonzero((np.isnan(a) != np.isnan(b)) | (np.isfinite(a) != np.isfinite(b)) | (np.isposinf(a) != np.isposinf(b)) | (np.isneginf(a) != np.isneginf(b))))
    differences = np.abs(a[both] - b[both])
    mismatches = int(np.count_nonzero(differences))
    return {"passed": mismatches == 0 and nonfinite_mismatches == 0,
            "shape": list(a.shape), "finite_values": int(both.sum()),
            "value_mismatches": mismatches, "nan_mask_mismatches": nonfinite_mismatches,
            "maximum_absolute_difference": float(differences.max(initial=0))}


def manifest_audit():
    manifest = readj(DELIVERY / "FILE_MANIFEST.json")
    entries = manifest["files"]
    json_paths = [e["path"] for e in entries]
    sha_entries = [line.split("  ", 1) for line in (DELIVERY / "MANIFEST.sha256").read_text(encoding="utf-8").splitlines() if line]
    sha_map = {p: h for h, p in sha_entries}
    all_files = {p.relative_to(DELIVERY).as_posix(): p for p in DELIVERY.rglob("*") if p.is_file()}
    expected_self = {"FILE_MANIFEST.json", "MANIFEST.sha256"}
    errors, evidence = [], []
    for e in entries:
        p = DELIVERY / e["path"]
        if not p.is_file():
            errors.append({"path": e["path"], "error": "missing"})
            continue
        actual = sha(p)
        row = {"path": e["path"], "bytes": p.stat().st_size, "sha256": actual,
               "json_sha_matches": actual == e["sha256"], "json_bytes_match": p.stat().st_size == e["bytes"],
               "sha_manifest_matches": sha_map.get(e["path"]) == actual}
        evidence.append(row)
        if not all(row[k] for k in ("json_sha_matches", "json_bytes_match", "sha_manifest_matches")):
            errors.append(row)
    unexpected = sorted(set(all_files) - set(json_paths) - expected_self)
    cache = [p for p in all_files if any(s in {"__pycache__", ".venv", "node_modules", ".pytest_cache", ".mypy_cache"} for s in Path(p).parts) or Path(p).suffix.lower() in {".pyc", ".pyo", ".tmp", ".partial"}]
    sensitive = [p for p in all_files if re.search(r"(^|/)(\.env(?:\..*)?|id_rsa|id_ed25519|credentials(?:\..*)?|secrets?(?:\..*)?|token(?:\..*)?)$|\.(pem|key|p12|pfx)$", p, re.I)]
    large = [{"path": p, "bytes": f.stat().st_size} for p, f in all_files.items() if f.stat().st_size > 90_000_000]
    report = {"passed": not errors and not unexpected and len(json_paths) == len(set(json_paths)) == manifest["file_count"] and set(json_paths) == set(sha_map),
              "source_commit_claim": manifest.get("source_commit"), "file_count_claim": manifest["file_count"],
              "json_entry_count": len(entries), "sha_entry_count": len(sha_entries), "total_files_including_manifests": len(all_files),
              "total_bytes": sum(p.stat().st_size for p in all_files.values()), "unexpected_unlisted_files": unexpected,
              "manifest_path_set_difference": sorted(set(json_paths) ^ set(sha_map)),
              "duplicate_json_entries": len(json_paths) - len(set(json_paths)), "duplicate_sha_entries": len(sha_entries)-len(sha_map),
              "files_over_90000000_bytes": large, "unexpected_cache_candidates": cache, "sensitive_filename_candidates": sensitive,
              "sensitive_scan_scope": "Filenames only. This is not an exhaustive secret-content scanner.", "errors": errors,
              "largest_files": sorted([{"path": p, "bytes": f.stat().st_size} for p, f in all_files.items()], key=lambda x: x["bytes"], reverse=True)[:8],
              "files": evidence}
    return report


def source_audit():
    pairs = [("inputs/problem.pdf", "A题/A题.pdf"), ("inputs/attachment1.xlsx", "A题/附件/附件1.xlsx"),
             ("inputs/attachment2.xlsx", "A题/附件/附件2.xlsx"),
             ("inputs/result2_template.xlsx", "A题/附件/附件3/result2.xlsx"),
             ("inputs/result4_template.xlsx", "A题/附件/附件3/result4.xlsx"),
             ("q123_closeout/inputs/attachment1.xlsx", "A题/附件/附件1.xlsx")]
    rows = []
    for delivered, original in pairs:
        a, b = sha(DELIVERY / delivered), sha(ROOT / original)
        rows.append({"delivery_path": delivered, "original_path": original, "delivery_sha256": a,
                     "original_sha256": b, "passed": a == b})
    return {"passed": all(x["passed"] for x in rows), "files": rows}


def csv_audit(path, expected, expected_header):
    with path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        header = next(reader)
        rows = [[np.nan if value == "" else float(value) for value in row] for row in reader]
    report = array_check(rows, expected)
    report["header_matches"] = header == expected_header
    report["passed"] &= report["header_matches"]
    report["data_rows"] = len(rows)
    return report


def q4_audit():
    z = dict(np.load(DELIVERY / "output/main.npz", allow_pickle=False))
    d = readj(DELIVERY / "output/result4_data.json")
    endpoint = readj(DELIVERY / "output/end_event.json")
    times = z["time_s"]
    keep = times > 0
    checks = {"json_time": array_check(d["time_s"], times[keep]), "json_radius": array_check(d["radius_m"], z["radius_m"][keep]),
              "json_C_fixed": array_check(d["C_fixed"], z["sample_TC"][keep, :, 1]),
              "json_C_surface": array_check(d["C_surface"], z["surface_TC"][keep, 1])}
    grid_expected = np.r_[np.arange(0, endpoint["execution_s"], 60.0), endpoint["execution_s"]]
    checks["time_grid"] = array_check(times, grid_expected)
    attachment = openpyxl.load_workbook(DELIVERY / "inputs/attachment2.xlsx", read_only=True, data_only=True)
    raw_radius = np.array([[row[0], row[1]] for row in list(attachment.active.values)[1:] if isinstance(row[0], (float, int))], dtype=float)
    attachment.close()
    radius_expected = np.interp(times, raw_radius[:, 0], raw_radius[:, 1] / 100)
    radius_diff = float(np.max(np.abs(radius_expected-z["radius_m"])))
    checks["radius_from_attachment2_linear_interpolation"] = {"passed": radius_diff < 1e-14, "maximum_absolute_difference_m": radius_diff,
        "source_rows": len(raw_radius), "source_time_range_s": [float(raw_radius[0, 0]), float(raw_radius[-1, 0])], "policy": "Linear interpolation; constant endpoint if outside measurement time span."}
    outside = z["fixed_radius_m"][None, :] > z["radius_m"][:, None] + 1e-12
    checks["npz_domain_mask"] = {"passed": bool(np.array_equal(np.isnan(z["sample_TC"][:, :, 0]), outside) and np.array_equal(np.isnan(z["sample_TC"][:, :, 1]), outside)), "outside_positions_including_initial_state": int(outside.sum())}
    trajectory_expected = np.column_stack([times, z["radius_m"], z["max_C"], z["max_radius_m"], z["mean_C"], z["surface_TC"], z["sample_TC"][:, :, 0], z["sample_TC"][:, :, 1]])
    trajectory_header = ["time_s", "radius_m", "max_C", "max_radius_m", "mean_C", "T_surface_C", "C_surface"] + [f"T_r{j/10:.1f}cm_C" for j in range(21)] + [f"C_r{j/10:.1f}cm" for j in range(21)]
    checks["trajectory_csv"] = csv_audit(DELIVERY / "output/trajectory_60s_unrounded.csv", trajectory_expected, trajectory_header)
    table_times = np.r_[np.arange(21600, endpoint["execution_s"], 21600.0), endpoint["execution_s"]]
    indexes = [int(np.where(times == t)[0][0]) for t in table_times]
    table_expected = np.column_stack([table_times / 3600, z["sample_TC"][indexes, ::5, 1], z["surface_TC"][indexes, 1], z["radius_m"][indexes] * 100])
    checks["table6_csv"] = csv_audit(DELIVERY / "output/table6_unrounded.csv", table_expected, ["time_h", "r0cm", "r0.5cm", "r1cm", "r1.5cm", "r2cm", "surface", "radius_cm"])
    md_rows = [line.strip().strip("|").split("|") for line in (DELIVERY / "output/表6.md").read_text(encoding="utf-8").splitlines() if line.startswith("|")][2:]
    expected_strings = [["—" if not np.isfinite(v) else f"{v:.4f}" for v in row] for row in table_expected]
    checks["table6_markdown"] = {"passed": md_rows == expected_strings, "data_rows": len(md_rows), "format": "Each finite value displayed to four decimal places; out-of-domain em dash."}
    workbook = openpyxl.load_workbook(DELIVERY / "output/result4.xlsx", read_only=True, data_only=False)
    sheet = workbook["Sheet1"]
    notes = {1: "输出说明", 2: "时间单位：秒；固定径向坐标单位：厘米。", 3: "水分浓度为干基含水率，单位kg水/kg干物质。",
             4: "固定径向点超出当前半径时留空，不代表零。", 5: "表面值取r=R(t)；与固定点重合时两列数值一致。",
             6: "含水率按十进制四位ROUND_HALF_UP舍入；未舍入轨迹另附，严格终判使用未舍入全域最大值。"}
    expected_header = ["时间\\到药材中心的距离", *[j/10 for j in range(21)], "药材表面", None, "实际半径R(t) / cm", None, notes[1]]
    errors, numeric_count, blanks, examined, n, radius_error, coincident = [], 0, 0, 0, 0, 0.0, 0
    for row_number, row in enumerate(sheet.iter_rows(values_only=False), 1):
        values = [cell.value for cell in row]
        examined += len(values)
        if len(values) > 27 and any(v is not None for v in values[27:]):
            errors.append([row_number, "Unexpected cells beyond AA"])
        values = (values + [None] * 27)[:27]
        if any(cell.data_type in {"f", "e"} for cell in row):
            errors.append([row_number, "Formula or error cell"])
        if row_number == 1:
            if values != expected_header:
                errors.append([row_number, "Header mismatch"])
            continue
        n += 1
        if n >= len(times):
            errors.append([row_number, "Unexpected trailing row"])
            continue
        expected = [float(times[n]), *[r4(v) for v in z["sample_TC"][n, :, 1]], r4(z["surface_TC"][n, 1]), None, float(z["radius_m"][n]*100), None, notes.get(row_number)]
        for col, (got, want) in enumerate(zip(values, expected), 1):
            if col == 25 and isinstance(got, (int, float)):
                radius_error = max(radius_error, abs(got-want))
                if abs(got-want) <= 1e-12:
                    continue
            if got != want:
                errors.append([row_number, col, str(got), str(want)])
        numeric_count += sum(v is not None for v in values[1:23])
        blanks += sum(v is None for v in values[1:22])
        for j, r in enumerate(z["fixed_radius_m"]):
            if abs(r - z["radius_m"][n]) < 1e-12:
                coincident += 1
                if values[j+1] != values[22]:
                    errors.append([row_number, j+2, "surface mismatch"])
    workbook.close()
    checks["result4_xlsx"] = {"passed": not errors and n == len(times)-1 and workbook.sheetnames == ["Sheet1"], "sheet_names": workbook.sheetnames,
            "all_returned_cells_examined": examined, "data_rows": n, "numeric_moisture_cells_compared": numeric_count,
            "outside_domain_blank_cells": blanks, "surface_coincidences_checked": coincident, "radius_maximum_difference_cm": radius_error,
            "error_count": len(errors), "first_100_errors": errors[:100], "first_time_s": float(times[1]), "last_time_s": float(times[-1]),
            "coverage": "Every row, all returned cells including core data A:W, blank spacers X/Z, radius Y and notes AA. No Excel application execution or new visual render."}
    return {"passed": all(c["passed"] for c in checks.values()), "checks": checks,
            "array_shapes": {k: list(v.shape) for k, v in z.items()}, "source_sha256": {p: sha(DELIVERY / p) for p in ["output/main.npz", "output/result4.xlsx", "output/result4_data.json"]}}


def parse_cell(cell, strings):
    kind = cell.get("t")
    if cell.find(NS+"f") is not None or kind == "e":
        return ("ERROR_OR_FORMULA", kind)
    if kind == "inlineStr":
        return "".join(cell.itertext())
    v = cell.find(NS+"v")
    if v is None:
        return None
    if kind == "s":
        return strings[int(v.text)]
    if kind == "str":
        return v.text
    if kind in {None, "n"}:
        return float(v.text)
    return ("UNEXPECTED_TYPE", kind, v.text)


def q2_audit():
    z = dict(np.load(DELIVERY / "q123_closeout/output/q23_unified.npz", allow_pickle=False))
    times = z["time_s"]
    rounded = round_array(z["profile_TC"][1:])
    path = DELIVERY / "q123_closeout/result2.xlsx"
    reports = {}
    with zipfile.ZipFile(path) as archive:
        relns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
        wb = ET.fromstring(archive.read("xl/workbook.xml"))
        rel = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {entry.get("Id"): entry.get("Target") for entry in rel}
        sheets = [(entry.get("name"), targets[entry.get(relns+"id")]) for entry in wb.find(NS+"sheets")]
        strings = []
        if "xl/sharedStrings.xml" in archive.namelist():
            strings = ["".join(s.itertext()) for s in ET.fromstring(archive.read("xl/sharedStrings.xml"))]
        for component, (name, target) in enumerate(sheets):
            xml_path = "xl/" + target if not target.startswith("/") else target.lstrip("/")
            rows, count, time_errors, value_errors, coordinate_errors, formula_errors, maxdiff, first_errors = 0, 0, 0, 0, 0, 0, 0.0, []
            header = None
            with archive.open(xml_path) as stream:
                for _, element in ET.iterparse(stream, events=("end",), tag=NS+"row"):
                    rownum = int(element.get("r"))
                    cells = element.findall(NS+"c")
                    vals = [parse_cell(c, strings) for c in cells]
                    if rownum == 1:
                        header = vals
                    else:
                        rows += 1
                        expected_coords = [f"{chr(65+j)}{rownum}" for j in range(22)]
                        coordinate_errors += int([c.get("r") for c in cells] != expected_coords or rownum != rows+1)
                        if rows >= len(times) or len(vals) != 22 or any(not isinstance(v, (int, float)) for v in vals):
                            formula_errors += 1
                            if len(first_errors) < 10:
                                first_errors.append({"row": rownum, "reason": "count/type/overflow"})
                        else:
                            time_errors += vals[0] != float(times[rows])
                            diff = np.abs(np.asarray(vals[1:], dtype=float) - rounded[rows-1, :, component])
                            value_errors += int(np.count_nonzero(diff))
                            maxdiff = max(maxdiff, float(diff.max(initial=0)))
                            count += len(vals)-1
                            if diff.any() and len(first_errors)<10:
                                first_errors.append({"row": rownum, "max_abs_difference": float(diff.max())})
                    element.clear()
                    while element.getprevious() is not None:
                        del element.getparent()[0]
            header_ok = header == ["时间\\到药材中心的距离", *z["radius_cm"].tolist()]
            reports[name] = {"passed": rows == len(times)-1 and not any([time_errors, value_errors, coordinate_errors, formula_errors]) and header_ok,
                "data_rows": rows, "numeric_profile_cells_compared": count, "header_matches": header_ok,
                "time_mismatches": int(time_errors), "value_mismatches": value_errors, "row_or_cell_coordinate_mismatches": coordinate_errors,
                "type_or_count_errors": formula_errors, "maximum_absolute_difference_after_rounding": maxdiff, "first_errors": first_errors,
                "first_time_s": float(times[1]), "last_time_s": float(times[-1])}
            print(json.dumps({"stage": "q2_sheet_complete", "sheet": name, **reports[name]}, ensure_ascii=False), flush=True)
    time_grid_passed = bool(np.array_equal(times[:-1], np.arange(len(times)-1, dtype=float)) and times[-2] < times[-1] < times[-2]+1)
    return {"passed": all(x["passed"] for x in reports.values()) and [x[0] for x in sheets] == ["温度", "水分浓度"] and time_grid_passed,
        "sheets": reports, "workbook_sha256": sha(path), "npz_sha256": sha(DELIVERY / "q123_closeout/output/q23_unified.npz"),
        "time_grid_passed": time_grid_passed,
        "coverage": "All 206907 data rows in both sheets and every one of 8690094 profile values; all time values, headers, row/cell coordinates and formula/error types. Compared against saved NPZ after independent ROUND_HALF_UP rounding; no PDE solve, Excel recalculation or visual render.",
        "rounding": "Independent NumPy implementation away from ties; Decimal ROUND_HALF_UP for near ties (within 1e-8 of scaled half)."}


def main():
    started = time.perf_counter()
    report = {"audit_utc": datetime.now(timezone.utc).isoformat(), "root": str(ROOT), "delivery": str(DELIVERY),
              "runtime": {"python": sys.version, "executable": sys.executable, "numpy": np.__version__, "openpyxl": openpyxl.__version__, "platform": platform.platform()},
              "scope": "Independent read-only file, source identity, saved-array, CSV and complete workbook consistency audit. No numerical solve and no source workbook edits."}
    initial_manifests = {p: sha(DELIVERY/p) for p in ["MANIFEST.sha256", "FILE_MANIFEST.json"]}
    report["manifest"] = manifest_audit()
    print(json.dumps({"stage": "manifest", "passed": report["manifest"]["passed"], "files": report["manifest"]["total_files_including_manifests"]}), flush=True)
    report["original_inputs"] = source_audit()
    report["q4"] = q4_audit()
    print(json.dumps({"stage": "q4", "passed": report["q4"]["passed"]}), flush=True)
    report["q2"] = q2_audit()
    report["manifest_files_unchanged_during_audit"] = initial_manifests == {p: sha(DELIVERY/p) for p in initial_manifests}
    report["passed"] = all(report[k]["passed"] for k in ["manifest", "original_inputs", "q4", "q2"]) and report["manifest_files_unchanged_during_audit"]
    report["elapsed_s"] = time.perf_counter() - started
    (HERE / "artifact_audit.json").write_text(json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False)+"\n", encoding="utf-8")
    m, q4, q2 = report["manifest"], report["q4"]["checks"], report["q2"]
    lines = ["# 第四问交付文件与表格独立核查", "", f"审计状态：{'通过' if report['passed'] else '存在不一致，需查看JSON'}。核查日期：{report['audit_utc']}。", "",
             f"交付目录共 {m['total_files_including_manifests']} 个文件，其中两个清单以外的 {m['json_entry_count']} 个条目已逐个重算 SHA-256 与字节数；两份清单一致，未列文件 {len(m['unexpected_unlisted_files'])} 个。",
             f"原题PDF、附件1/2、result2/4模板与前三问收尾使用的附件1共6项，全部与项目 A题 原件逐字节哈希相同。",
             f"超过90,000,000字节的文件 {len(m['files_over_90000000_bytes'])} 个；缓存候选 {len(m['unexpected_cache_candidates'])} 个；敏感文件名候选 {len(m['sensitive_filename_candidates'])} 个。敏感检查仅针对文件名，不等同完整密钥内容扫描。", "",
             f"result4.xlsx 全部 {q4['result4_xlsx']['data_rows']} 个数据行已核对，涵盖 {q4['result4_xlsx']['numeric_moisture_cells_compared']} 个含水率数值、{q4['result4_xlsx']['outside_domain_blank_cells']} 个域外空格、全部时间和半径、间隔列、说明文字。固定位置与表面重合的 {q4['result4_xlsx']['surface_coincidences_checked']} 格已核对。",
             "result4_data.json 的所有数组与 main.npz 一致；60秒时间网格加实际结束点、附件2线性半径、NPZ域外掩码、完整CSV轨迹、表6 CSV以及表6 Markdown全部相符。",
             f"完整 result2.xlsx 两张表各 {next(iter(q2['sheets'].values()))['data_rows']} 个数据行，合计 {sum(s['numeric_profile_cells_compared'] for s in q2['sheets'].values()):,} 个剖面数值全部流式核对；时间、坐标连续性、表头、值类型、四位十进制舍入均相符。末时刻206906.76秒。", "",
             "这些核查证明交付文件完整以及工作簿/CSV/已保存数组彼此一致，不是对PDE、材料参数、物理假设或实验预测的独立验证。本审计没有重新运行求解器，没有修改交付原包，没有在Excel应用中重算，也没有新渲染工作簿。", "",
             "复用方式：在仓库根目录以安装了NumPy、openpyxl、lxml的Python执行 `python review_q4_20260912/artifacts/audit_artifacts.py`。详细机器证据见 `artifact_audit.json`；脚本仅向本审计目录写输出。", ""]
    if not report["passed"]:
        lines = ["# 第四问交付文件与表格独立核查", "", "审计未通过。以下状态来自机器报告，不能将未通过的项目视为已经一致。", ""]
        lines.extend([f"- {k}: {'通过' if report[k]['passed'] else '失败，详见 artifact_audit.json'}" for k in ["manifest", "original_inputs", "q4", "q2"]])
        lines.extend(["", "已核查范围与具体差异均保存在 artifact_audit.json。失败可能源于检查器或交付内容，需要具体定位，不能仅依据总状态认定交付数据错误。", ""])
    if (HERE / "first_run_checker_issue.json").exists():
        lines.extend(["首轮检查器未支持 OOXML 的 `t=str` 普通文本格，曾将Q2表头误判不一致；修正后重新执行了全量核对。失败记录及归因保存在 `first_run_checker_issue.json`，属于检查器问题，未改动交付文件。", ""])
    (HERE / "artifact_audit.md").write_text("\n".join(lines), encoding="utf-8")
    print(json.dumps({"stage": "complete", "passed": report["passed"], "elapsed_s": report["elapsed_s"]}), flush=True)


if __name__ == "__main__":
    main()
