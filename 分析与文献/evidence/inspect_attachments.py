"""Read-only source audit. Does not save or alter supplied workbooks."""
from pathlib import Path
from collections import Counter
import hashlib
import json
import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "A题" / "附件"
OUT = Path(__file__).resolve().parent


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fit_saturating(t, y):
    """Descriptive fit only: y_inf + (y0-y_inf)*exp(-(t/tau)^p)."""
    initial = float(y[0])
    def predict(theta, tx):
        asymptote, tau, exponent = theta
        return asymptote + (initial - asymptote) * np.exp(-(tx / tau) ** exponent)

    # Two-dimensional deterministic grid refinement; for each tau,p the
    # amplitude has an exact least-squares solution. No external optimizer.
    log_bounds = [np.log(60.), np.log(30000.)]
    p_bounds = [0.3, 3.]
    for refinement in range(6):
        logs = np.linspace(*log_bounds, 51)
        exponents = np.linspace(*p_bounds, 51)
        taus = np.exp(logs)
        basis = 1 - np.exp(-(t[None, None, :] / taus[None, :, None]) ** exponents[:, None, None])
        amp = np.sum(basis * (y - initial), axis=-1) / np.sum(basis ** 2, axis=-1)
        sse = np.sum((initial + amp[:, :, None] * basis - y) ** 2, axis=-1)
        i, j = np.unravel_index(np.argmin(sse), sse.shape)
        theta = np.array([initial + amp[i, j], taus[j], exponents[i]])
        dl, dp = logs[1] - logs[0], exponents[1] - exponents[0]
        log_bounds = [logs[j] - 2 * dl, logs[j] + 2 * dl]
        p_bounds = [max(0.01, exponents[i] - 2 * dp), exponents[i] + 2 * dp]
    residual = predict(theta, t) - y
    return {
        "form": "y_inf + (y0-y_inf)*exp(-(t/tau)^p); t in seconds",
        "initial_value_fixed_to_source": initial,
        "parameters": dict(zip(["y_inf", "tau_seconds", "p"], theta.tolist())),
        "fit_method": "six deterministic refinements of 51x51 tau,p grid, analytic least-squares amplitude",
        "rmse": float(np.sqrt(np.mean(residual ** 2))),
        "mae": float(np.mean(np.abs(residual))),
        "max_abs_error": float(np.max(np.abs(residual))),
        "r_squared": float(1 - np.sum(residual ** 2) / np.sum((y - np.mean(y)) ** 2)),
        "predicted_at_48h": float(predict(theta, np.array([172800.]))[0]),
        "predicted_at_72h": float(predict(theta, np.array([259200.]))[0]),
    }, predict, theta


def numeric_stats(values):
    a = np.asarray(values, dtype=float)
    return {
        "count": len(a), "min": float(a.min()), "max": float(a.max()),
        "mean": float(a.mean()), "sample_sd": float(a.std(ddof=1)),
        "first": float(a[0]), "last": float(a[-1]),
        "increasing_steps": int(np.sum(np.diff(a) > 0)),
        "decreasing_steps": int(np.sum(np.diff(a) < 0)),
        "equal_steps": int(np.sum(np.diff(a) == 0)),
    }


audit = {"scope": "read-only structural and numerical audit, no PDE solution", "files": []}
arrays = {}
for path in sorted(SOURCE.rglob("*.xlsx")):
    if path.name.startswith("~$"):
        continue  # Office session lock metadata is not an input workbook.
    before = digest(path)
    wb = load_workbook(path, read_only=True, data_only=False)
    entry = {"path": str(path.relative_to(ROOT)), "sha256": before, "sheets": []}
    for sheet in wb.worksheets:
        rows = list(sheet.values)
        populated = [(ri + 1, ci + 1, value) for ri, row in enumerate(rows)
                     for ci, value in enumerate(row) if value is not None]
        numeric_body = [row for row in rows[1:] if isinstance(row[0], (int, float))]
        entry["sheets"].append({
            "name": sheet.title,
            "declared_range": sheet.calculate_dimension(),
            "true_nonempty_range": f"{get_column_letter(min(c for r,c,v in populated))}{min(r for r,c,v in populated)}:{get_column_letter(max(c for r,c,v in populated))}{max(r for r,c,v in populated)}",
            "nonempty_cells": len(populated),
            "nonempty_rows": len(set(r for r, c, v in populated)),
            "header": list(rows[0]),
            "first_rows": [list(row) for row in rows[:5]],
            "last_rows": [list(row) for row in rows[-3:]],
            "formula_cells": [f"{get_column_letter(c)}{r}" for r, c, v in populated if isinstance(v, str) and v.startswith("=")],
            "error_cells": [f"{get_column_letter(c)}{r}" for r, c, v in populated if isinstance(v, str) and v.startswith("#")],
            "numeric_time_rows": len(numeric_body),
            "body_result_nonempty_cells": sum(v is not None for row in rows[1:] for v in row[1:]),
        })
        if path.name in {"附件1.xlsx", "附件2.xlsx"}:
            data = np.asarray(rows[1:], dtype=float)
            arrays[path.name] = data
            intervals = np.diff(data[:, 0])
            entry["data_audit"] = {
                "record_count": len(data),
                "time_min_seconds": float(data[0, 0]),
                "time_max_seconds": float(data[-1, 0]),
                "time_max_hours": float(data[-1, 0] / 3600),
                "time_interval_seconds_counts": dict(Counter(str(float(x)) for x in intervals)),
                "duplicate_times": len(data) - len(set(data[:, 0])),
                "time_strictly_increasing": bool(np.all(intervals > 0)),
                "missing_values_in_data_rectangle": sum(v is None for row in rows[1:] for v in row),
                "nonfinite_values": int(np.sum(~np.isfinite(data))),
                "negative_values": int(np.sum(data < 0)),
                "fields": {str(rows[0][col]): numeric_stats(data[:, col]) for col in range(1, data.shape[1])},
            }
    wb.close()
    entry["source_hash_unchanged_after_read"] = digest(path) == before
    audit["files"].append(entry)

env = arrays["附件1.xlsx"]
t = env[:, 0]
audit["environment_analysis"] = {}
for col, name, unit in [(1, "temperature", "degree_C"), (2, "moisture", "kg/kg")]:
    y = env[:, col]
    fit, predict, theta = fit_saturating(t, y)
    train = t <= 10800
    holdout_fit, holdout_predict, holdout_theta = fit_saturating(t[train], y[train])
    residual_holdout = holdout_predict(holdout_theta, t[~train]) - y[~train]
    tail = t >= 10800
    tail_slope = np.polyfit(t[tail] / 3600, y[tail], 1)[0]
    audit["environment_analysis"][name] = {
        "unit": unit, "full_fit": fit,
        "holdout_train_range_s": [0, 10800],
        "holdout_test_range_s": [10860, 14400],
        "holdout_rmse": float(np.sqrt(np.mean(residual_holdout ** 2))),
        "last_hour_source_range": f"Sheet1!{get_column_letter(col+1)}182:{get_column_letter(col+1)}242",
        "last_hour_stats": numeric_stats(y[tail]),
        "last_hour_linear_slope_per_hour": float(tail_slope),
        "last_hour_mean_extrapolation": float(np.mean(y[tail])),
        "last_observation_extrapolation": float(y[-1]),
    }
radius = arrays["附件2.xlsx"]
audit["radius_analysis"] = {
    "radius_units": "cm", "time_units": "s",
    "relative_radius_loss_at_72h": float(1-radius[-1, 1]/radius[0, 1]),
    "cross_section_area_ratio_at_72h": float((radius[-1, 1]/radius[0, 1])**2),
    "every_6h": radius[::12].tolist(),
    "first_radius_below_1_5cm": radius[np.flatnonzero(radius[:, 1] < 1.5)[0]].tolist(),
    "monotone_nonincreasing": bool(np.all(np.diff(radius[:, 1]) <= 0)),
    "last_24h_radius_range": [float(radius[radius[:, 0] >= 172800, 1].min()), float(radius[radius[:, 0] >= 172800, 1].max())],
}
OUT.mkdir(parents=True, exist_ok=True)
(OUT / "attachment_audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({"environment_analysis": audit["environment_analysis"], "radius_analysis": audit["radius_analysis"],
                  "source_hashes_unchanged": all(f["source_hash_unchanged_after_read"] for f in audit["files"]),
                  "saved": str(OUT / "attachment_audit.json")}, ensure_ascii=False, indent=2))
